import argparse
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX


def normalize(value):
    return "" if value is None else str(value).strip()


def require_columns(sheet, labels):
    headers = [normalize(cell.value) for cell in sheet[1]]
    lookup = {header.casefold(): index + 1 for index, header in enumerate(headers) if header}
    missing = [label for label in labels if label.casefold() not in lookup]
    if missing:
        raise ValueError(
            f"Sheet '{sheet.title}' is missing columns: {', '.join(missing)}. "
            f"Found: {', '.join(headers)}"
        )
    return lookup


def validate_id(value, label, sheet_name, row_number):
    if not value:
        raise ValueError(f"Sheet '{sheet_name}', row {row_number}: {label} is blank")
    if len(value) != 4 or not value.isalnum():
        raise ValueError(
            f"Sheet '{sheet_name}', row {row_number}: {label} '{value}' "
            "must contain exactly four letters or numbers"
        )


def read_employees(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Employee sheet '{sheet_name}' was not found")
    sheet = workbook[sheet_name]
    lookup = require_columns(sheet, ["Employee ID", "Employee Name", "Manager"])
    employees = {}
    for row_number in range(2, sheet.max_row + 1):
        employee_id = normalize(sheet.cell(row_number, lookup["employee id"]).value).upper()
        employee_name = normalize(sheet.cell(row_number, lookup["employee name"]).value)
        manager_value = normalize(sheet.cell(row_number, lookup["manager"]).value)
        if not employee_id and not employee_name and not manager_value:
            continue
        validate_id(employee_id, "employee ID", sheet_name, row_number)
        if not employee_name:
            raise ValueError(f"Sheet '{sheet_name}', row {row_number}: employee name is blank")
        if employee_id in employees:
            raise ValueError(f"Sheet '{sheet_name}' has duplicate employee ID '{employee_id}'")
        employees[employee_id] = {
            "id": employee_id,
            "name": employee_name,
            "isManager": manager_value.casefold() in {
                "☑ yes", "yes", "true", "1", "x", "checked", "manager"
            },
            "sourceRow": row_number,
        }
    if not employees:
        raise ValueError(f"Sheet '{sheet_name}' does not contain any employees")
    return employees


def theme_rgb(workbook, theme_index):
    if not workbook.loaded_theme:
        return None
    root = ET.fromstring(workbook.loaded_theme)
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", namespace)
    if scheme is None:
        return None
    colors = {}
    for element in list(scheme):
        key = element.tag.rsplit("}", 1)[-1]
        child = next(iter(element), None)
        if child is not None:
            value = child.attrib.get("lastClr") or child.attrib.get("val")
            if value and len(value) >= 6:
                colors[key] = value[-6:].upper()
    order = [
        "lt1", "dk1", "lt2", "dk2", "accent1", "accent2",
        "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink",
    ]
    return colors.get(order[theme_index]) if 0 <= theme_index < len(order) else None


def apply_tint(rgb, tint):
    tint = float(tint or 0)
    channels = [int(rgb[index:index + 2], 16) for index in (0, 2, 4)]
    adjusted = []
    for channel in channels:
        value = channel * (1 + tint) if tint < 0 else channel * (1 - tint) + 255 * tint
        adjusted.append(max(0, min(255, round(value))))
    return "".join(f"{channel:02X}" for channel in adjusted)


def cell_fill_hex(workbook, cell):
    if cell.fill.fill_type != "solid":
        raise ValueError(
            f"Sheet '{cell.parent.title}', cell {cell.coordinate}: choose a solid fill color"
        )
    color = cell.fill.fgColor
    rgb = None
    if color.type == "rgb" and color.rgb:
        rgb = color.rgb[-6:].upper()
    elif color.type == "indexed" and color.indexed is not None:
        index = int(color.indexed)
        if 0 <= index < len(COLOR_INDEX):
            rgb = COLOR_INDEX[index][-6:].upper()
    elif color.type == "theme" and color.theme is not None:
        rgb = theme_rgb(workbook, int(color.theme))
    if not rgb:
        raise ValueError(
            f"Sheet '{cell.parent.title}', cell {cell.coordinate}: the fill color could not be read. "
            "Use a standard solid Excel fill color."
        )
    return f"#{apply_tint(rgb, color.tint)}"


def read_role_colors(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Role settings sheet '{sheet_name}' was not found")
    sheet = workbook[sheet_name]
    lookup = require_columns(sheet, ["Role Type", "Slide Color"])
    colors = {}
    for row_number in range(2, sheet.max_row + 1):
        role = normalize(sheet.cell(row_number, lookup["role type"]).value)
        if not role:
            continue
        key = role.casefold()
        if key in colors:
            raise ValueError(f"Sheet '{sheet_name}' has duplicate role type '{role}'")
        colors[key] = {
            "name": role,
            "color": cell_fill_hex(workbook, sheet.cell(row_number, lookup["slide color"])),
        }
    if not colors:
        raise ValueError(f"Sheet '{sheet_name}' does not contain any role types")
    return colors


def read_state_sheet(workbook, sheet_name, columns, employees, role_colors):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' was not found")
    sheet = workbook[sheet_name]
    pcn_label = columns.get("pcn", "PCN")
    explanation_label = columns.get("explanation", "Explanation")
    labels = [columns[key] for key in ("employeeId", "managerId", "role")] + [pcn_label, explanation_label]
    lookup = require_columns(sheet, labels)
    records = []
    seen = {}
    for row_number in range(2, sheet.max_row + 1):
        employee_id = normalize(sheet.cell(row_number, lookup[columns["employeeId"].casefold()]).value).upper()
        manager_id = normalize(sheet.cell(row_number, lookup[columns["managerId"].casefold()]).value).upper()
        role = normalize(sheet.cell(row_number, lookup[columns["role"].casefold()]).value)
        pcn = normalize(sheet.cell(row_number, lookup[pcn_label.casefold()]).value)
        explanation = normalize(sheet.cell(row_number, lookup[explanation_label.casefold()]).value)
        if not employee_id and not pcn and not manager_id and not role and not explanation:
            continue
        if employee_id:
            validate_id(employee_id, "employee ID", sheet_name, row_number)
        elif not pcn:
            raise ValueError(f"Sheet '{sheet_name}', row {row_number}: enter an employee or a PCN for the vacant position")
        validate_id(manager_id, "manager ID", sheet_name, row_number)
        if employee_id and employee_id not in employees:
            raise ValueError(f"Sheet '{sheet_name}', row {row_number}: employee ID '{employee_id}' is not on Employees")
        if manager_id not in employees:
            raise ValueError(f"Sheet '{sheet_name}', row {row_number}: manager ID '{manager_id}' is not on Employees")
        if not employees[manager_id]["isManager"]:
            raise ValueError(f"Sheet '{sheet_name}', row {row_number}: '{manager_id}' is not marked as a manager")
        if not role:
            role = "Unspecified"
        if role.casefold() not in role_colors:
            raise ValueError(f"Sheet '{sheet_name}', row {row_number}: role '{role}' is not on Role Settings")
        record_key = f"EMP:{employee_id}" if employee_id else f"PCN:{pcn.upper()}"
        if record_key in seen:
            duplicate_label = f"employee ID '{employee_id}'" if employee_id else f"PCN '{pcn}'"
            raise ValueError(
                f"Sheet '{sheet_name}' has duplicate {duplicate_label} "
                f"on rows {seen[record_key]} and {row_number}"
            )
        seen[record_key] = row_number
        records.append({
            "employeeId": employee_id,
            "employeeName": employees[employee_id]["name"] if employee_id else "Vacant",
            "managerId": manager_id,
            "managerName": employees[manager_id]["name"],
            "role": role_colors[role.casefold()]["name"],
            "pcn": pcn,
            "explanation": explanation,
            "isVacant": not bool(employee_id),
            "sourceRow": row_number,
        })
    return records


def main():
    parser = argparse.ArgumentParser(description="Extract organization workbook data to JSON")
    parser.add_argument("--settings", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    settings_path = Path(args.settings).resolve()
    settings = json.loads(settings_path.read_text(encoding="utf-8"))
    workbook_path = Path(settings["workbookPath"])
    if not workbook_path.is_absolute():
        workbook_path = (settings_path.parent / workbook_path).resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    employees = read_employees(workbook, settings["sheets"].get("employees", "Employees"))
    role_colors = read_role_colors(workbook, settings["sheets"].get("roles", "Role Settings"))
    payload = {
        "workbookPath": str(workbook_path),
        "roleColors": {item["name"]: item["color"] for item in role_colors.values()},
        "employees": list(employees.values()),
        "current": read_state_sheet(workbook, settings["sheets"]["current"], settings["columns"], employees, role_colors),
        "future": read_state_sheet(workbook, settings["sheets"]["future"], settings["columns"], employees, role_colors),
    }
    Path(args.output).write_text(json.dumps(payload, indent=2), encoding="utf-8")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
